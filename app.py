import sys
import os
import subprocess
import threading
import platform
import unicodedata
import re
import sqlite3
import time
from datetime import datetime
from pathlib import Path
from copy import copy

import requests

from PySide6.QtWidgets import (
    QApplication,
    QWidget,
    QLabel,
    QLineEdit,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QHBoxLayout,
    QMessageBox,
    QProgressBar,
    QDialog,
    QTextEdit,
    QListWidget,
    QListWidgetItem,
    QSplitter,
    QFrame,
)
from PySide6.QtCore import QEvent, QTimer, Qt, QStandardPaths

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

# --- IMPORTAR EL ACTUALIZADOR DESDE TU OTRO ARCHIVO ---
from actualizador import VERSION_ACTUAL, MonitorActualizaciones
from metricas import PanelMetricas




# ============================================================

# CONFIGURACIÓN

# ============================================================



BASE_URL = "https://dco.dducaimunixela.org"



LOGIN_URL = BASE_URL + "/login_2.php"



PRINCIPAL_URL = BASE_URL + "/principal_2.php"



BUSQUEDA_URL = (

    BASE_URL

    + "/area_restringida/consultar_expediente/"

    + "buscar_opciones.php"

)



# ============================================================

# EDITAR EXPEDIENTE FINALIZADO

# ============================================================



EDITAR_EXP_FINALIZADO_URL = (

    BASE_URL

    + "/AR_editar_exp_finalizado.php"

)



BUSQUEDA_EDITAR_FINALIZADO_URL = (

    BASE_URL

    + "/area_restringida/editar_exp_finalizado/"

    + "buscar_opciones2.php"

)



DATOS_EDITAR_FINALIZADO_URL = (

    BASE_URL

    + "/area_restringida/editar_exp_finalizado/"

    + "obtener_datos_seleccionados.php"

)





# ============================================================

# BÚSQUEDA GLOBAL DEL EXCEL EN EL SISTEMA

# ============================================================



NOMBRE_ARCHIVO_EXCEL = "03_FORMATO DE IMPRESION DE LICENCIAS 2026.xlsm"

RUTA_EXCEL_CACHE = None





def obtener_ruta_excel():

    global RUTA_EXCEL_CACHE



    if RUTA_EXCEL_CACHE and os.path.exists(RUTA_EXCEL_CACHE):

        return RUTA_EXCEL_CACHE



    home = Path.home()



    ubicaciones_frecuentes = [

        home / "Desktop" / NOMBRE_ARCHIVO_EXCEL,

        home / "OneDrive" / "Desktop" / NOMBRE_ARCHIVO_EXCEL,

        home / "Escritorio" / NOMBRE_ARCHIVO_EXCEL,

        home / "Documents" / NOMBRE_ARCHIVO_EXCEL,

        home / "Documentos" / NOMBRE_ARCHIVO_EXCEL,

        home / "Downloads" / NOMBRE_ARCHIVO_EXCEL,

        home / "Descargas" / NOMBRE_ARCHIVO_EXCEL,

        home / NOMBRE_ARCHIVO_EXCEL,

    ]



    for ruta in ubicaciones_frecuentes:

        if ruta.exists():

            RUTA_EXCEL_CACHE = str(ruta)

            return RUTA_EXCEL_CACHE



    carpetas_excluidas = {

        "Library",

        "System",

        "Applications",

        ".Trash",

        "AppData",

        "Program Files",

        "Program Files (x86)",

        "Windows",

        "node_modules",

        ".git",

    }



    try:

        for raíz, directorios, archivos in os.walk(home):

            directorios[:] = [

                d for d in directorios if d not in carpetas_excluidas and not d.startswith(".")

            ]

            if NOMBRE_ARCHIVO_EXCEL in archivos:

                ruta_encontrada = os.path.join(raíz, NOMBRE_ARCHIVO_EXCEL)

                RUTA_EXCEL_CACHE = ruta_encontrada

                return RUTA_EXCEL_CACHE

    except Exception:

        pass



    return None





def limpiar_texto(texto):

    if not texto:

        return ""

    texto = str(texto).strip().upper()

    texto = unicodedata.normalize('NFD', texto)

    texto = ''.join(c for c in texto if unicodedata.category(c) != 'Mn')

    texto = re.sub(r'\s+', ' ', texto)

    return texto





def formatear_fecha(fecha_str):

    """

    Convierte una fecha formato YYYY-MM-DD a DD/MM/YYYY.

    Si viene con horas (ej: 2025-07-01 14:30:00), las elimina.

    """

    if not fecha_str:

        return ""

    fecha_str = str(fecha_str).strip().split()[0]

    if "-" in fecha_str:

        partes = fecha_str.split("-")

        if len(partes) == 3 and len(partes[0]) == 4:

            return f"{partes[2]}/{partes[1]}/{partes[0]}"

    return fecha_str





# ============================================================

# EVENTOS

# ============================================================



class LoginEvent(QEvent):



    EVENT_TYPE = QEvent.Type(

        QEvent.registerEventType()

    )



    def __init__(

        self,

        exito,

        mensaje

    ):



        super().__init__(

            self.EVENT_TYPE

        )



        self.exito = exito

        self.mensaje = mensaje





class SearchEvent(QEvent):



    EVENT_TYPE = QEvent.Type(

        QEvent.registerEventType()

    )



    def __init__(

        self,

        resultados,

        mensaje=""

    ):



        super().__init__(

            self.EVENT_TYPE

        )



        self.resultados = resultados

        self.mensaje = mensaje





# ============================================================

# FUNCIONES GENERALES

# ============================================================



def valor(

    datos,

    clave

):



    if not datos:



        return ""



    resultado = datos.get(

        clave,

        ""

    )



    if resultado is None:



        return ""



    return str(

        resultado

    ).strip()





# ============================================================

# ORDENAR NOMBRES

# ============================================================



def ordenar_nombre(

    nombre

):



    if not nombre:



        return ""



    nombre = str(

        nombre

    ).strip()



    if not nombre:



        return ""



    nombre = nombre.strip(

        " \t\n\r,.;:"

    )



    if not nombre:



        return ""



    marcas_condueño = {

        "CDNO",

        "CDÑO",

        "CDÑA",

        "CDNA",

        "CONDUEÑO",

        "CONDUEÑA",

        "CONDUEÑOS",

        "CONDUEÑAS",

    }



    palabras = nombre.split()



    if not palabras:



        return ""



    palabras_limpias = []



    for palabra in palabras:



        palabra = palabra.strip(

            " \t\n\r,.;:"

        )



        if palabra:



            palabras_limpias.append(

                palabra

            )



    palabras = palabras_limpias



    if not palabras:



        return ""



    condueño = []



    while palabras:



        ultima = palabras[-1].upper().strip(

            " \t\n\r,.;:"

        )



        if ultima in marcas_condueño:



            condueño.insert(

                0,

                palabras.pop()

            )



        else:



            break



    if palabras:



        ultima = palabras[-1].upper().strip(

            " \t\n\r,.;:"

        )



        if ultima == "Y" and condueño:



            condueño.insert(

                0,

                palabras.pop()

            )



    if not palabras:



        return " ".join(

            condueño

        ).strip(

            " \t\n\r,.;:"

        )



    if len(palabras) == 1:



        resultado = palabras[0]



        if condueño:



            resultado += (

                " "

                + " ".join(condueño)

            )



        return resultado.strip(

            " \t\n\r,.;:"

        )



    palabras_mayus = [

        p.upper()

        for p in palabras

    ]



    cantidad_apellidos = 2



    if palabras_mayus[0] == "DE":



        cantidad_apellidos = 3



        if (

            len(palabras) > 1

            and palabras_mayus[1]

            in {

                "LA",

                "LOS",

                "LAS",

            }

        ):



            cantidad_apellidos = 4



    if "Y" in palabras_mayus[:5]:



        posicion_y = (

            palabras_mayus.index(

                "Y"

            )

        )



        if posicion_y >= 1:



            cantidad_apellidos = max(

                cantidad_apellidos,

                posicion_y + 2

            )



    for i in range(

        min(

            5,

            len(palabras_mayus) - 2

        )

    ):



        if (

            palabras_mayus[i] == "DE"

            and palabras_mayus[i + 1]

            in {

                "LA",

                "LOS",

                "LAS",

            }

        ):



            cantidad_apellidos = max(

                cantidad_apellidos,

                i + 4

            )



    cantidad_apellidos = min(

        cantidad_apellidos,

        len(palabras) - 1

    )



    apellidos = palabras[

        :cantidad_apellidos

    ]



    nombres = palabras[

        cantidad_apellidos:

    ]



    resultado = nombres + apellidos



    texto_final = " ".join(

        resultado

    )



    if condueño:



        texto_final += (

            " "

            + " ".join(condueño)

        )



    texto_final = texto_final.strip(

        " \t\n\r,.;:"

    )



    texto_final = " ".join(

        texto_final.split()

    )



    return texto_final





# ============================================================

# PROFESIONAL

# ============================================================



def identificar_profesional(

    nombre

):



    if not nombre:



        return "", ""



    nombre = str(

        nombre

    ).strip()



    if not nombre:



        return "", ""



    nombre = nombre.strip(

        " \t\n\r,.;:"

    )



    partes = nombre.split()



    if not partes:



        return "", ""



    primera = partes[0].upper().strip(

        " \t\n\r,.;:"

    )



    if primera == "ING":



        nombre_limpio = " ".join(

            partes[1:]

        )



        nombre_limpio = nombre_limpio.strip(

            " \t\n\r,.;:"

        )



        return "ING", nombre_limpio



    if primera == "ARQ":



        nombre_limpio = " ".join(

            partes[1:]

        )



        nombre_limpio = nombre_limpio.strip(

            " \t\n\r,.;:"

        )



        return "ARQ", nombre_limpio



    return "", nombre





# ============================================================

# EXTRAER RESULTADOS

# ============================================================



def extraer_resultados(

    respuesta

):



    if isinstance(

        respuesta,

        list

    ):



        return respuesta



    if isinstance(

        respuesta,

        dict

    ):



        posibles = [

            "data",

            "datos",

            "resultados",

            "results",

            "expedientes",

        ]



        for clave in posibles:



            valor_encontrado = (

                respuesta.get(

                    clave

                )

            )



            if isinstance(

                valor_encontrado,

                list

            ):



                return valor_encontrado



    return []





# ============================================================

# EXTRAER CAMPOS DE EDITAR EXPEDIENTE FINALIZADO

# ============================================================



def extraer_campos_editar_finalizado(

    datos

):



    if not datos:



        return {}



    if isinstance(

        datos,

        list

    ):



        if not datos:



            return {}



        datos = datos[0]



    if not isinstance(

        datos,

        dict

    ):



        return {}



    campos_necesarios = [

        "AREA_DEL_TERRENO",

        "AREA_TERRENO",

        "area_terreno",

        "area_del_predio",

        "T_METROS",

        "MURO",

        "V_ALINEACION",

        "RECIBO_ALINEACION",

        "V_LICENCIA_SUELO",

        "RECIBO_LICENCIA_SUELO",

        "VIGENCIA_EN_MESES",

        "SE_VENCE",

        "VISITAS_MINIMAS",

        # Número de colegiado del profesional planificador
        "colegiado_planificado",
        "COLEGIADO_PLANIFICADO",
        "colegiado_planificador",
        "COLEGIADO_PLANIFICADOR",

    ]



    resultado = {}



    for campo in campos_necesarios:



        valor_campo = datos.get(

            campo,

            ""

        )



        if valor_campo is None:



            valor_campo = ""



        resultado[campo] = str(

            valor_campo

        ).strip()



    return resultado





# ============================================================

# MOSTRAR ARCHIVO EN EXPLORADOR

# ============================================================



def mostrar_archivo_en_explorador(

    ruta

):



    if not ruta:



        return False



    ruta_absoluta = os.path.abspath(

        ruta

    )



    if not os.path.exists(

        ruta_absoluta

    ):



        return False



    sistema = platform.system()



    try:



        if sistema == "Windows":



            subprocess.run(

                [

                    "explorer.exe",

                    f'/select,"{ruta_absoluta}"'

                ],

                check=False

            )



        elif sistema == "Darwin":



            subprocess.run(

                [

                    "open",

                    "-R",

                    ruta_absoluta

                ],

                check=False

            )



        elif sistema == "Linux":



            subprocess.run(

                [

                    "xdg-open",

                    os.path.dirname(

                        ruta_absoluta

                    )

                ],

                check=False

            )



        else:



            return False



        return True



    except Exception:



        return False





# ============================================================


# ============================================================
# RECONEXIÓN AUTOMÁTICA E HISTORIAL LOCAL
# ============================================================

INTERVALO_CONEXION_MS = 3 * 60 * 1000


class ConexionEvent(QEvent):
    EVENT_TYPE = QEvent.Type(QEvent.registerEventType())

    def __init__(self, ok, texto, monitor=False):
        super().__init__(self.EVENT_TYPE)
        self.ok = ok
        self.texto = texto
        self.monitor = monitor


class HistorialExportaciones:
    """Historial privado por perfil del sistema operativo; no guarda contraseñas."""

    COLUMNAS = (
        "fecha", "expediente", "propietario", "direccion", "zona",
        "tramite", "licencia", "licencia_suelo", "planificador", "ejecutor",
        "colegiado", "archivo_excel", "hoja", "fila_excel",
    )

    def __init__(self):
        ubicacion = QStandardPaths.writableLocation(
            QStandardPaths.StandardLocation.AppDataLocation
        )
        if not ubicacion:
            ubicacion = str(Path.home() / ".dco_expedientes")
        carpeta = Path(ubicacion)
        carpeta.mkdir(parents=True, exist_ok=True)
        if os.name != "nt":
            try:
                carpeta.chmod(0o700)
            except OSError:
                pass
        self.ruta = carpeta / "historial_exportaciones.sqlite3"
        self._crear_tabla()
        if os.name != "nt":
            try:
                self.ruta.chmod(0o600)
            except OSError:
                pass

    def _abrir(self):
        con = sqlite3.connect(str(self.ruta), timeout=15)
        con.row_factory = sqlite3.Row
        return con

    def _crear_tabla(self):
        with self._abrir() as con:
            con.execute("""
                CREATE TABLE IF NOT EXISTS exportaciones (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    fecha TEXT NOT NULL, expediente TEXT NOT NULL,
                    propietario TEXT, direccion TEXT, zona TEXT, tramite TEXT,
                    licencia TEXT, licencia_suelo TEXT, planificador TEXT,
                    ejecutor TEXT, colegiado TEXT, archivo_excel TEXT,
                    hoja TEXT, fila_excel INTEGER
                )
            """)
            con.execute(
                "CREATE INDEX IF NOT EXISTS ix_historial_fecha "
                "ON exportaciones(fecha DESC)"
            )

    def guardar(self, registro):
        datos = {col: registro.get(col, "") for col in self.COLUMNAS}
        datos["fecha"] = datetime.now().astimezone().isoformat(timespec="seconds")
        columnas = ", ".join(self.COLUMNAS)
        marcadores = ", ".join("?" for _ in self.COLUMNAS)
        with self._abrir() as con:
            con.execute(
                f"INSERT INTO exportaciones ({columnas}) VALUES ({marcadores})",
                tuple(datos[col] for col in self.COLUMNAS),
            )

    def listar(self, texto="", limite=250):
        filtro = f"%{texto.strip()}%"
        with self._abrir() as con:
            filas = con.execute(
                """SELECT id, fecha, expediente, propietario, licencia
                   FROM exportaciones
                   WHERE expediente LIKE ? OR propietario LIKE ? OR licencia LIKE ?
                   ORDER BY id DESC LIMIT ?""",
                (filtro, filtro, filtro, limite),
            ).fetchall()
            return [dict(fila) for fila in filas]

    def obtener(self, identificador):
        with self._abrir() as con:
            fila = con.execute(
                "SELECT * FROM exportaciones WHERE id=?", (identificador,)
            ).fetchone()
            return dict(fila) if fila else None


# APLICACIÓN

# ============================================================



class DCOApp(QWidget):

    def __init__(self):
        super().__init__()
        self.setWindowTitle(f"DCO - Consulta de Expedientes V{VERSION_ACTUAL}")
        self.resize(1330, 740)

        self.session = requests.Session()
        self._session_lock = threading.RLock()
        self._autenticado = False
        self._cerrando = False
        self._verificacion_en_curso = False
        self._login_activo = False
        self._busqueda_activa = False
        self._excel_activo = False
        self.resultados = []
        # Credenciales de la cuenta autorizada usada por esta instalación.
        # IMPORTANTE: cualquier persona con acceso al código o ejecutable
        # podría recuperarlas. No distribuir a usuarios no autorizados.
        self.usuario = "jose_estrada"
        self.password = "OLi842cYKYM0k3Q07i"
        self.historial = HistorialExportaciones()
        self.panel_metricas = PanelMetricas(self.historial.ruta)

        self.crear_interfaz()
        self.actualizar_panel_historial()
        self.monitor_actualizaciones = MonitorActualizaciones(
            parent=self, puede_notificar=self.puede_mostrar_actualizacion
        )

        # Monitor de sesión separado del monitor de actualizaciones de GitHub.
        self.temporizador_conexion = QTimer(self)
        self.temporizador_conexion.setInterval(INTERVALO_CONEXION_MS)
        self.temporizador_conexion.timeout.connect(self.verificar_conexion)
        self.temporizador_conexion.start()
        QTimer.singleShot(0, self.iniciar_sesion_automatica)

    def puede_mostrar_actualizacion(self):
        """La ventana de actualización solo aparece si no hay tareas en curso."""
        return not any((
            self._login_activo,
            self._busqueda_activa,
            self._excel_activo,
        ))


    def comprobar_actualizaciones_manualmente(self):
        """Botón opcional para comprobar la versión sin reiniciar DCO."""
        self.monitor_actualizaciones.comprobar(manual=True)


    # ========================================================

    # INTERFAZ

    # ========================================================



    def crear_interfaz(

        self

    ):



        layout_principal = QVBoxLayout()
        layout_principal.setContentsMargins(12, 10, 12, 10)



        # ----------------------------------------------------

        # BÚSQUEDA

        # ----------------------------------------------------



        fila_busqueda = QHBoxLayout()



        self.expediente_input = QLineEdit()



        self.expediente_input.setPlaceholderText(

            "Ingrese número de expediente"

        )



        self.expediente_input.returnPressed.connect(

            self.buscar

        )



        self.buscar_button = QPushButton(

            "BUSCAR"

        )



        self.buscar_button.clicked.connect(

            self.buscar

        )



        self.buscar_button.setEnabled(False)



        fila_busqueda.addWidget(

            QLabel("Expediente:")

        )



        fila_busqueda.addWidget(

            self.expediente_input

        )



        fila_busqueda.addWidget(

            self.buscar_button

        )



        layout_principal.addLayout(

            fila_busqueda

        )

        # Indicador de meta diaria: datos del historial local de exportaciones.
        layout_principal.addWidget(self.panel_metricas)



        # ----------------------------------------------------

        # TABLA

        # ----------------------------------------------------



        self.tabla = QTableWidget()



        self.tabla.setColumnCount(

            7

        )



        self.tabla.setHorizontalHeaderLabels(

            [

                "EXPEDIENTE",

                "PROPIETARIO",

                "DIRECCIÓN",

                "ZONA",

                "TRÁMITE",

                "LICENCIA",

                "LICENCIA SUELO",

            ]

        )



        self.tabla.setEditTriggers(

            QTableWidget.EditTrigger.NoEditTriggers

        )



        self.tabla.setSelectionBehavior(

            QTableWidget.SelectionBehavior.SelectRows

        )



        self.tabla.doubleClicked.connect(

            self.ver_detalle

        )



        layout_principal.addWidget(

            self.tabla

        )



        # ----------------------------------------------------

        # BOTONES

        # ----------------------------------------------------



        fila_botones = QHBoxLayout()



        self.detalle_button = QPushButton(

            "VER DETALLE"

        )



        self.detalle_button.clicked.connect(

            self.ver_detalle

        )



        self.excel_button = QPushButton(

            "EXPORTAR A EXCEL"

        )



        self.excel_button.clicked.connect(

            self.exportar_excel

        )



        fila_botones.addWidget(

            self.detalle_button

        )



        fila_botones.addWidget(

            self.excel_button

        )

        # Permite solicitar una revisión sin cerrar la aplicación.
        self.actualizar_button = QPushButton("BUSCAR ACTUALIZACIONES")
        self.actualizar_button.clicked.connect(
            self.comprobar_actualizaciones_manualmente
        )
        fila_botones.addWidget(self.actualizar_button)

        layout_principal.addLayout(

            fila_botones

        )



        # ----------------------------------------------------

        # PROGRESO

        # ----------------------------------------------------



        self.progreso = QProgressBar()



        self.progreso.setRange(

            0,

            100

        )



        self.progreso.setValue(

            0

        )



        layout_principal.addWidget(

            self.progreso

        )



        # ----------------------------------------------------

        # ESTADO

        # ----------------------------------------------------



        self.estado = QLabel(

            "Conectando al sistema..."

        )



        layout_principal.addWidget(

            self.estado

        )



        zona_trabajo = QWidget()
        zona_trabajo.setLayout(layout_principal)

        divisor = QSplitter(Qt.Orientation.Horizontal, self)
        divisor.addWidget(self.crear_panel_izquierdo())
        divisor.addWidget(zona_trabajo)
        divisor.setChildrenCollapsible(False)
        divisor.setSizes([310, 1020])
        contenedor = QVBoxLayout(self)
        contenedor.setContentsMargins(4, 4, 4, 4)
        contenedor.addWidget(divisor)






    def crear_panel_izquierdo(self):
        panel = QFrame()
        panel.setObjectName("PanelHistorial")
        panel.setMinimumWidth(245)
        panel.setMaximumWidth(480)
        panel.setStyleSheet("""
            QFrame#PanelHistorial { background-color: palette(base);
                                    border-right: 1px solid palette(mid); }
            QListWidget { border: 1px solid palette(mid); border-radius: 5px; }
        """)
        caja = QVBoxLayout(panel)
        caja.setContentsMargins(10, 10, 10, 10)
        titulo = QLabel("HISTORIAL DE EXPORTACIONES")
        titulo.setStyleSheet("font-weight: 700; font-size: 13px;")
        titulo.setWordWrap(True)
        caja.addWidget(titulo)
        self.conexion_estado = QLabel("● Conectando...")
        self.conexion_estado.setStyleSheet("color: #bb8700; font-weight: 600;")
        caja.addWidget(self.conexion_estado)
        self.reconexion_button = QPushButton("RECONECTAR AL PORTAL")
        self.reconexion_button.clicked.connect(self.reconectar_manualmente)
        caja.addWidget(self.reconexion_button)
        self.historial_busqueda = QLineEdit()
        self.historial_busqueda.setPlaceholderText("Buscar expediente, propietario o licencia")
        self.historial_busqueda.textChanged.connect(self.actualizar_panel_historial)
        caja.addWidget(self.historial_busqueda)
        self.historial_total = QLabel("0 registros recientes")
        caja.addWidget(self.historial_total)
        self.historial_lista = QListWidget()
        self.historial_lista.itemSelectionChanged.connect(self.mostrar_detalle_historial)
        self.historial_lista.itemDoubleClicked.connect(self.copiar_expediente_historial)
        caja.addWidget(self.historial_lista, 3)
        caja.addWidget(QLabel("DATOS GUARDADOS"))
        self.historial_detalle = QTextEdit()
        self.historial_detalle.setReadOnly(True)
        self.historial_detalle.setMinimumHeight(155)
        caja.addWidget(self.historial_detalle, 2)
        btn = QPushButton("USAR EXPEDIENTE SELECCIONADO")
        btn.clicked.connect(self.copiar_expediente_historial)
        caja.addWidget(btn)
        btn_archivo = QPushButton("LOCALIZAR EXCEL")
        btn_archivo.clicked.connect(self.localizar_excel_historial)
        caja.addWidget(btn_archivo)
        nota = QLabel("Historial local de exportaciones confirmadas. No modifica Excel.")
        nota.setWordWrap(True)
        nota.setStyleSheet("color: palette(mid);")
        caja.addWidget(nota)
        return panel

    def actualizar_panel_historial(self, *_):
        texto = self.historial_busqueda.text().strip()
        try:
            registros = self.historial.listar(texto)
        except Exception as exc:
            self.historial_total.setText(f"Error leyendo el historial: {exc}")
            return
        self.historial_lista.clear()
        for reg in registros:
            fecha = reg["fecha"].replace("T", " ")[:16]
            nombre = reg.get("propietario", "")
            item = QListWidgetItem(
                f"{reg['expediente']}  ·  {fecha}\n{nombre}"
            )
            item.setData(Qt.ItemDataRole.UserRole, reg["id"])
            self.historial_lista.addItem(item)
        self.historial_total.setText(f"{len(registros)} registros (máximo 250)")
        self.historial_detalle.clear()

    def registro_historial_seleccionado(self):
        item = self.historial_lista.currentItem()
        return self.historial.obtener(item.data(Qt.ItemDataRole.UserRole)) if item else None

    def mostrar_detalle_historial(self):
        reg = self.registro_historial_seleccionado()
        if not reg:
            self.historial_detalle.clear()
            return
        campos = [
            ("Exportado", reg["fecha"].replace("T", " ")[:19]),
            ("Expediente", reg["expediente"]),
            ("Propietario", reg["propietario"]),
            ("Dirección", reg["direccion"]),
            ("Zona", reg["zona"]),
            ("Trámite", reg["tramite"]),
            ("Licencia", reg["licencia"]),
            ("Licencia suelo", reg["licencia_suelo"]),
            ("Planificador", reg["planificador"]),
            ("Ejecutor", reg["ejecutor"]),
            ("Colegiado", reg["colegiado"]),
            ("Hoja / fila", f"{reg['hoja']} / {reg['fila_excel']}"),
            ("Excel", reg["archivo_excel"]),
        ]
        self.historial_detalle.setPlainText(
            "\n".join(f"{nombre}: {dato}" for nombre, dato in campos if dato)
        )

    def copiar_expediente_historial(self, *_):
        reg = self.registro_historial_seleccionado()
        if not reg:
            return
        self.expediente_input.setText(reg["expediente"])
        self.expediente_input.setFocus()
        self.estado.setText("Expediente recuperado del historial; pulsa BUSCAR para consultarlo.")

    def localizar_excel_historial(self):
        reg = self.registro_historial_seleccionado()
        if reg and reg.get("archivo_excel") and os.path.isfile(reg["archivo_excel"]):
            mostrar_archivo_en_explorador(reg["archivo_excel"])
        else:
            QMessageBox.information(
                self, "Historial", "Selecciona una exportación cuyo archivo Excel todavía exista."
            )

    def indicar_conexion(self, ok, mensaje=""):
        self._autenticado = bool(ok)
        self.conexion_estado.setText(
            "● Portal conectado" if ok else "● Portal desconectado (reintentando)"
        )
        self.conexion_estado.setStyleSheet(
            "color: #16834c; font-weight: 600;" if ok else
            "color: #b13a39; font-weight: 600;"
        )
        self.buscar_button.setEnabled(ok and not self._busqueda_activa)
        if mensaje:
            self.estado.setText(mensaje)

    def reconectar_manualmente(self):
        """Reintenta el acceso con la cuenta integrada; no solicita credenciales."""
        if self._login_activo or self._verificacion_en_curso:
            return
        self.iniciar_sesion_automatica()

    def _autenticar_con_bloqueo(self):
        """Solo llamar mientras se mantiene self._session_lock."""
        if not (self.usuario and self.password):
            raise RuntimeError("No se han proporcionado las credenciales del portal.")
        nueva_sesion = requests.Session()
        r = nueva_sesion.get(LOGIN_URL, timeout=(10, 25))
        r.raise_for_status()
        respuesta = nueva_sesion.post(
            LOGIN_URL,
            data={"username": self.usuario, "password": self.password},
            allow_redirects=True,
            timeout=(10, 25),
        )
        respuesta.raise_for_status()
        if "principal_2.php" not in respuesta.url:
            nueva_sesion.close()
            raise RuntimeError("El portal rechazó el inicio de sesión. Revisa tus credenciales.")
        anterior = self.session
        self.session = nueva_sesion
        anterior.close()
        self._autenticado = True

    @staticmethod
    def _respuesta_requiere_login(respuesta):
        url_final = respuesta.url.lower().split("?", 1)[0]
        if url_final.endswith("/login_2.php") or respuesta.status_code in (401, 403):
            return True
        comienzo = respuesta.text[:1600].lower().lstrip() if respuesta.content else ""
        return (comienzo.startswith("<") and "username" in comienzo and
                "password" in comienzo and "login" in comienzo)

    def _solicitud_portal(self, metodo, url, *, requiere_json=False, **kwargs):
        """Serializa requests.Session y renueva credenciales si caducó el acceso."""
        with self._session_lock:
            if not self._autenticado:
                self._autenticar_con_bloqueo()
                QApplication.postEvent(
                    self, ConexionEvent(True, "Sesión del portal recuperada automáticamente.")
                )
            kwargs.setdefault("timeout", (10, 30))
            kwargs.setdefault("allow_redirects", True)
            try:
                respuesta = self.session.request(metodo, url, **kwargs)
                if self._respuesta_requiere_login(respuesta):
                    self._autenticado = False
                    self._autenticar_con_bloqueo()
                    respuesta = self.session.request(metodo, url, **kwargs)
                    if self._respuesta_requiere_login(respuesta):
                        raise RuntimeError("El portal sigue solicitando inicio de sesión.")
                    QApplication.postEvent(
                        self, ConexionEvent(True, "Se recuperó automáticamente la sesión vencida.")
                    )
                respuesta.raise_for_status()
                if requiere_json:
                    try:
                        respuesta.json()
                    except ValueError as exc:
                        raise RuntimeError(
                            "El portal devolvió una respuesta no válida (no es JSON)."
                        ) from exc
                return respuesta
            except requests.RequestException:
                self._autenticado = False
                QApplication.postEvent(
                    self, ConexionEvent(False, "Sin conexión con el portal; se reintentará automáticamente.")
                )
                raise

    def verificar_conexion(self):
        """Comprobación no bloqueante, cada 3 minutos y sin invadir tareas activas."""
        if (self._cerrando or self._login_activo or self._busqueda_activa or
                self._excel_activo or self._verificacion_en_curso or
                not (self.usuario and self.password)):
            return
        self._verificacion_en_curso = True
        threading.Thread(target=self._verificar_sesion_hilo, daemon=True).start()

    def _verificar_sesion_hilo(self):
        try:
            with self._session_lock:
                respuesta = None
                if self._autenticado:
                    respuesta = self.session.get(PRINCIPAL_URL, timeout=(8, 15))
                if (respuesta is None or self._respuesta_requiere_login(respuesta) or
                        "principal_2.php" not in respuesta.url):
                    self._autenticado = False
                    self._autenticar_con_bloqueo()
                    texto = "Sesión recuperada automáticamente."
                else:
                    respuesta.raise_for_status()
                    texto = "Portal conectado."
            QApplication.postEvent(self, ConexionEvent(True, texto, monitor=True))
        except Exception as exc:
            QApplication.postEvent(self, ConexionEvent(False, f"Portal no disponible: {exc}", monitor=True))

    # ========================================================
    # INICIO DE SESIÓN EN SEGUNDO PLANO
    # ========================================================

    def iniciar_sesion_automatica(self):
        if self._login_activo:
            return
        self._login_activo = True
        self.buscar_button.setEnabled(False)
        self.estado.setText("Conectando con el portal DCO...")
        self.progreso.setValue(20)
        threading.Thread(target=self.login_hilo, daemon=True).start()

    def login_hilo(self):
        try:
            with self._session_lock:
                self._autenticar_con_bloqueo()
            QApplication.postEvent(
                self, LoginEvent(True, "Sesión del portal iniciada correctamente.")
            )
        except Exception as exc:
            QApplication.postEvent(self, LoginEvent(False, f"No se pudo conectar: {exc}"))

    # ========================================================

    # BÚSQUEDA

    # ========================================================



    def buscar(

        self

    ):



        expediente = (

            self.expediente_input.text()

            .strip()

        )



        if not expediente:



            QMessageBox.warning(

                self,

                "Falta expediente",

                "Ingrese un número de expediente."

            )



            return



        if not self._autenticado:
            QMessageBox.warning(
                self, "Portal sin conexión",
                "Aún no se ha iniciado sesión. Pulsa RECONECTAR AL PORTAL."
            )
            return



        self._busqueda_activa = True
        self.buscar_button.setEnabled(

            False

        )



        self.estado.setText(

            "Buscando expediente..."

        )



        self.progreso.setValue(

            40

        )



        hilo = threading.Thread(

            target=self.buscar_hilo,

            args=(expediente,),

            daemon=True

        )



        hilo.start()





    # ========================================================

    # BÚSQUEDA HILO

    # ========================================================



    def buscar_hilo(

        self,

        expediente

    ):



        try:



            datos_busqueda = {

                "propietario": "",

                "expediente": expediente,

                "direccion": "",

                "codigo_de_inmueble": "",

                "finca": "",

                "folio": "",

                "libro": "",

                "No_LICENCIA_OBRA": "",

                "No_LICENCIA_SUELO": "",

            }



            respuesta = (

                self._solicitud_portal(
                    "POST", BUSQUEDA_URL,
                    json=datos_busqueda, requiere_json=True,
                )

            )



            respuesta.raise_for_status()



            try:



                datos = (

                    respuesta.json()

                )



            except ValueError as exc:
                raise RuntimeError("El servidor devolvió JSON incorrecto.") from exc



            resultados = (

                extraer_resultados(

                    datos

                )

            )



            QApplication.postEvent(

                self,

                SearchEvent(

                    resultados,

                    "Búsqueda terminada."

                )

            )



        except Exception as e:



            QApplication.postEvent(

                self,

                SearchEvent(

                    [],

                    f"Error en búsqueda: {e}"

                )

            )





    # ========================================================

    # EVENTOS

    # ========================================================



    def event(

        self,

        event

    ):



        if (

            event.type()

            == LoginEvent.EVENT_TYPE

        ):

            self._login_activo = False
            self.progreso.setValue(

                100 if event.exito else 0

            )



            self.estado.setText(

                event.mensaje

            )



            self.indicar_conexion(event.exito, event.mensaje)
            if not event.exito:
                QMessageBox.warning(
                    self, "Conexión al portal",
                    event.mensaje + "\n\nPuedes pulsar RECONECTAR AL PORTAL."
                )



            return True





        if event.type() == ConexionEvent.EVENT_TYPE:
            if event.monitor:
                self._verificacion_en_curso = False
            # Evitar reemplazar el mensaje de una exportación activa.
            mensaje = event.texto if not (self._excel_activo or self._busqueda_activa) else ""
            self.indicar_conexion(event.ok, mensaje)
            return True

        if (

            event.type()

            == SearchEvent.EVENT_TYPE

        ):



            # ------------------------------------------------

            # EXPORTACIÓN CORRECTA

            # ------------------------------------------------



            if (

                event.mensaje

                == "EXPORTACION_EXCEL_OK"

            ):

                self._excel_activo = False
                if isinstance(event.resultados, dict):
                    self.actualizar_panel_historial()
                    self.panel_metricas.refrescar()
                self.excel_button.setEnabled(

                    True

                )



                self.progreso.setValue(

                    100

                )



                self.estado.setText(

                    "Expediente procesado correctamente."

                )



                ruta_final = obtener_ruta_excel()



                QMessageBox.information(

                    self,

                    "Excel",

                    f"El expediente fue procesado correctamente "

                    f"en {NOMBRE_ARCHIVO_EXCEL}."

                )



                if ruta_final:
                    mostrar_archivo_en_explorador(ruta_final)
                if isinstance(event.resultados, dict) and event.resultados.get("historial_error"):
                    QMessageBox.warning(
                        self, "Historial local",
                        "El Excel se guardó correctamente, pero no se pudo guardar "
                        "el registro en el panel izquierdo:\n" +
                        event.resultados["historial_error"],
                    )
                return True





            # ------------------------------------------------

            # ERROR EXPORTACIÓN

            # ------------------------------------------------



            if event.mensaje.startswith(

                "EXPORTACION_EXCEL_ERROR:"

            ):

                self._excel_activo = False
                self.excel_button.setEnabled(

                    True

                )



                self.progreso.setValue(

                    0

                )



                error = event.mensaje[

                    len(

                        "EXPORTACION_EXCEL_ERROR:"

                    ):

                ]



                self.estado.setText(

                    "Error al procesar expediente."

                )



                QMessageBox.critical(

                    self,

                    "Error Excel",

                    "No se pudo exportar:\n\n"

                    + error

                )



                return True





            # ------------------------------------------------

            # BÚSQUEDA NORMAL

            # ------------------------------------------------

            self._busqueda_activa = False
            self.buscar_button.setEnabled(self._autenticado)



            self.progreso.setValue(

                100

            )



            self.estado.setText(

                event.mensaje

            )



            if event.mensaje.startswith("Error en búsqueda:"):
                QMessageBox.warning(self, "Consulta al portal", event.mensaje)
            else:
                self.mostrar_resultados(event.resultados)



            return True





        return super().event(

            event

        )





    # ========================================================

    # MOSTRAR RESULTADOS

    # ========================================================



    def mostrar_resultados(

        self,

        resultados

    ):



        self.resultados = resultados



        self.tabla.setRowCount(

            0

        )



        if not resultados:



            QMessageBox.information(

                self,

                "Búsqueda",

                "No se encontraron resultados."

            )



            return



        self.tabla.setRowCount(

            len(resultados)

        )



        for fila, datos in enumerate(

            resultados

        ):



            expediente = valor(

                datos,

                "expediente"

            )



            propietario = valor(

                datos,

                "propietario"

            )



            direccion = valor(

                datos,

                "direccion"

            )



            zona = valor(

                datos,

                "zona"

            )



            tramite = valor(

                datos,

                "nombre_tramite"

            )



            licencia = valor(

                datos,

                "No_LICENCIA_OBRA"

            )



            licencia_suelo = valor(

                datos,

                "No_LICENCIA_SUELO"

            )



            valores = [

                expediente,

                propietario,

                direccion,

                zona,

                tramite,

                licencia,

                licencia_suelo,

            ]



            for columna, texto in enumerate(

                valores

            ):



                self.tabla.setItem(

                    fila,

                    columna,

                    QTableWidgetItem(

                        texto

                    )

                )



        self.tabla.resizeColumnsToContents()





    # ========================================================

    # OBTENER SELECCIONADO

    # ========================================================



    def obtener_seleccionado(

        self

    ):



        fila = (

            self.tabla.currentRow()

        )



        if fila < 0:



            QMessageBox.warning(

                self,

                "Selección",

                "Seleccione un expediente."

            )



            return None



        if fila >= len(

            self.resultados

        ):



            return None



        return self.resultados[

            fila

        ]





    # ========================================================

    # VER DETALLE

    # ========================================================



    def ver_detalle(

        self

    ):



        datos = (

            self.obtener_seleccionado()

        )



        if not datos:



            return



        dialogo = QDialog(

            self

        )



        dialogo.setWindowTitle(

            "Detalle del expediente"

        )



        dialogo.resize(

            900,

            700

        )



        layout = QVBoxLayout()



        texto = QTextEdit()



        texto.setReadOnly(

            True

        )



        contenido = ""



        for clave, valor_dato in datos.items():



            contenido += (

                f"{clave}: "

                f"{valor_dato}\n\n"

            )



        texto.setPlainText(

            contenido

        )



        layout.addWidget(

            texto

        )



        dialogo.setLayout(

            layout

        )



        dialogo.exec()





    # ========================================================

    # OBTENER DATOS DE EDITAR EXPEDIENTE FINALIZADO

    # ========================================================



    def obtener_datos_editar_finalizado(

        self,

        datos

    ):



        identificador = valor(

            datos,

            "id"

        )



        if not identificador:



            identificador = valor(

                datos,

                "id_tramite_anio"

            )



        if not identificador:



            expediente = valor(

                datos,

                "expediente"

            )



            if not expediente:



                return {}



            respuesta_busqueda = (

                self._solicitud_portal(
                    "POST", BUSQUEDA_EDITAR_FINALIZADO_URL,
                    requiere_json=True, json={

                        "columna": "expediente",

                        "valor": expediente,

                    },

                    timeout=30

                )

            )



            respuesta_busqueda.raise_for_status()



            resultados_busqueda = (

                respuesta_busqueda.json()

            )



            opciones = (

                resultados_busqueda.get(

                    "opciones2",

                    []

                )

            )



            if not opciones:



                return {}



            primera_opcion = opciones[0]



            if isinstance(

                primera_opcion,

                dict

            ):



                identificador = (

                    primera_opcion.get(

                        "id"

                    )

                )



            else:



                identificador = (

                    primera_opcion

                )



        if not identificador:



            return {}



        respuesta_datos = (

            self._solicitud_portal(
                "POST", DATOS_EDITAR_FINALIZADO_URL,
                requiere_json=True, json={

                    "id": identificador

                },

                timeout=30

            )

        )



        respuesta_datos.raise_for_status()



        datos_completos = (

            respuesta_datos.json()

        )



        return extraer_campos_editar_finalizado(

            datos_completos

        )





    # ========================================================

    # EXPORTAR A EXCEL

    # ========================================================



    def exportar_excel(

        self

    ):



        datos = (

            self.obtener_seleccionado()

        )



        if not datos:



            return



        self._excel_activo = True
        self.excel_button.setEnabled(

            False

        )



        self.progreso.setValue(

            20

        )



        self.estado.setText(

            "Buscando archivo y procesando datos..."

        )



        hilo = threading.Thread(

            target=self.exportar_excel_hilo,

            args=(datos,),

            daemon=True

        )



        hilo.start()





    # ========================================================

    # EXPORTAR A EXCEL HILO

    # ========================================================



    def exportar_excel_hilo(

        self,

        datos

    ):



        try:



            campos_adicionales = (

                self.obtener_datos_editar_finalizado(

                    datos

                )

            )



            datos_completos = dict(

                datos

            )



            # Mantener los datos originales cuando un campo adicional venga vacío.
            # Así no se pierde el colegiado recibido en la búsqueda principal.
            for clave, contenido in campos_adicionales.items():
                if contenido != "" or clave not in datos_completos:
                    datos_completos[clave] = contenido



            registro = self.crear_fila_excel(datos_completos)
            error_historial = ""
            try:
                self.historial.guardar(registro)
            except Exception as exc:
                error_historial = str(exc)
            QApplication.postEvent(
                self,
                SearchEvent(
                    {"registro": registro, "historial_error": error_historial},
                    "EXPORTACION_EXCEL_OK",
                )
            )



        except Exception as e:



            QApplication.postEvent(

                self,

                SearchEvent(

                    [],

                    f"EXPORTACION_EXCEL_ERROR:{e}"

                )

            )





    # ========================================================

    # ACTUALIZAR FILA EN LA PESTAÑA "BASE DE DATOS"

    # ========================================================



    def crear_fila_excel(

        self,

        datos

    ):



        expediente_valor = valor(

            datos,

            "expediente"

        )



        propietario_original = valor(

            datos,

            "propietario"

        )



        propietario = ordenar_nombre(

            propietario_original

        )



        # ----------------------------------------------------

        # DIRECCIÓN

        # ----------------------------------------------------



        direccion = valor(

            datos,

            "direccion"

        )



        zona = valor(

            datos,

            "zona"

        )



        if zona and zona.isdigit():



            direccion_excel = (

                f"{direccion} ZONA {zona}"

            )



        else:



            direccion_excel = direccion



        # ----------------------------------------------------

        # PROFESIONALES

        # ----------------------------------------------------



        tipo_planificador, planificador = (

            identificar_profesional(

                valor(

                    datos,

                    "planificador"

                )

            )

        )



        # Buscar el colegiado del planificador en los datos del expediente.
        # Se admiten variaciones frecuentes del nombre del campo.
        colegiado_planificado = (
            valor(datos, "colegiado_planificado")
            or valor(datos, "COLEGIADO_PLANIFICADO")
            or valor(datos, "colegiado_planificador")
            or valor(datos, "COLEGIADO_PLANIFICADOR")
        )

        tipo_ejecutor, ejecutor = (

            identificar_profesional(

                valor(

                    datos,

                    "ejecutor"

                )

            )

        )



        # ----------------------------------------------------

        # DATOS PRINCIPALES

        # ----------------------------------------------------



        licencia = valor(

            datos,

            "No_LICENCIA_OBRA"

        )



        finca = valor(

            datos,

            "finca"

        )



        folio = valor(

            datos,

            "folio"

        )



        libro = valor(

            datos,

            "libro"

        )



        codigo_municipal = valor(

            datos,

            "codigo_de_inmueble"

        )



        # ----------------------------------------------------

        # ÁREAS

        # ----------------------------------------------------



        area_del_terreno = (

            valor(datos, "AREA_DEL_TERRENO")

            or valor(datos, "AREA_TERRENO")

            or valor(datos, "area_terreno")

            or valor(datos, "METROS_TERRENO")

            or valor(datos, "area_del_predio")

        )



        area_a_construir = (

            valor(datos, "T_METROS")

            or valor(datos, "METROS")

            or valor(datos, "area_a_construir")

        )



        # ----------------------------------------------------

        # CAMPOS EDITAR FINALIZADO

        # ----------------------------------------------------



        muro = valor(

            datos,

            "MURO"

        )



        pago_alineacion = valor(

            datos,

            "V_ALINEACION"

        )



        recibo_alineacion = valor(

            datos,

            "RECIBO_ALINEACION"

        )



        pago_uso_suelo = valor(

            datos,

            "V_LICENCIA_SUELO"

        )



        recibo_uso_suelo = valor(

            datos,

            "RECIBO_LICENCIA_SUELO"

        )



        vigencia_en_meses = valor(

            datos,

            "VIGENCIA_EN_MESES"

        )



        # Se aplica la conversión de formato a las fechas

        se_vence = formatear_fecha(valor(datos, "SE_VENCE"))



        visitas_minimas = valor(

            datos,

            "VISITAS_MINIMAS"

        )



        # ----------------------------------------------------

        # RESTO DE DATOS

        # ----------------------------------------------------



        niveles = valor(

            datos,

            "NIVELES"

        )



        costo_obra = valor(

            datos,

            "COSTO_DE_OBRA"

        )



        destino_obra = valor(

            datos,

            "DESTINO_AUTORIZADO"

        )



        unidades_funcionales = valor(

            datos,

            "UNIDADES_FUNCIONALES_DE_ESTA_LICENCIA"

        )



        nombre_proyecto = valor(

            datos,

            "nombre_del_proyecto"

        )



        gestor = valor(

            datos,

            "gestor"

        )



        solicita = (

            valor(datos, "solicita")

            or valor(datos, "SOLICITA")

            or valor(datos, "nombre_tramite")

        )



        solicita_mayus = solicita.upper().strip(

            " \t\n\r,.;:"

        )



        numero_licencia_suelo = valor(

            datos,

            "No_LICENCIA_SUELO"

        )



        # Se aplica la conversión de formato a la fecha

        fecha_pago = formatear_fecha(valor(datos, "FECHA_PAGO"))



        # ====================================================

        # LOCALIZAR ARCHIVO EN LA COMPUTADORA

        # ====================================================



        excel_path = obtener_ruta_excel()



        if not excel_path:



            raise Exception(

                f"No se encontró el archivo '{NOMBRE_ARCHIVO_EXCEL}' en la computadora."

            )



        # ====================================================

        # CARGAR ARCHIVO PRESERVANDO MACROS Y FORMATOS

        # ====================================================



        wb = load_workbook(

            excel_path,

            keep_vba=True

        )



        # ====================================================

        # BUSCAR LA PESTAÑA "BASE DE DATOS"

        # ====================================================

        ws = None

        target_sheet_name = "BASE DE DATOS"



        for sheet_name in wb.sheetnames:

            if limpiar_texto(sheet_name) == target_sheet_name:

                ws = wb[sheet_name]

                break

        if not ws:

            raise Exception(

                f"No se encontró ninguna pestaña llamada 'Base de datos' en el archivo."

            )



        # ====================================================

        # MAPEADO DE COLUMNAS DE LA FILA 2

        # ====================================================



        fila_encabezados = 2

        mapa_columnas = {}



        for col in range(1, ws.max_column + 1):



            val = ws.cell(row=fila_encabezados, column=col).value



            if val is not None:



                enc_limpio = limpiar_texto(val)



                if enc_limpio not in mapa_columnas:



                    mapa_columnas[enc_limpio] = []



                mapa_columnas[enc_limpio].append(col)



        def obtener_columna(encabezado, ocurrencia=1):



            enc_limpio = limpiar_texto(encabezado)



            cols = mapa_columnas.get(enc_limpio, [])



            if len(cols) >= ocurrencia:



                return cols[ocurrencia - 1]



            return None



        columna_expediente = obtener_columna("EXPEDIENTE", 1)



        if not columna_expediente:

            raise Exception(

                f"En la pestaña '{ws.title}', no se encontró la columna 'EXPEDIENTE' en la fila {fila_encabezados}."

            )



        # ====================================================

        # COMPROBAR EXPEDIENTE EXISTENTE

        # ====================================================



        existe = False

        fila_existente = None



        for fila in range(fila_encabezados + 1, ws.max_row + 1):



            valor_existente = ws.cell(row=fila, column=columna_expediente).value



            if (

                valor_existente is not None

                and str(valor_existente).strip() != ""

                and str(valor_existente).strip() == expediente_valor

            ):



                existe = True

                fila_existente = fila



                break



        if existe:

            raise Exception(

                f"El expediente '{expediente_valor}' ya se encuentra registrado en la fila {fila_existente} "

                "de la 'BASE DE DATOS'. \n\nElimínalo primero de Excel y guarda el archivo si deseas volver a exportarlo."

            )



        # ====================================================

        # ENCONTRAR FILA VACÍA (TRATANDO 0, 0.0 Y FÓRMULAS COMO VACÍO)

        # ====================================================

        fila_destino = fila_encabezados + 1

        max_cols = ws.max_column

        while True:

            vacia = True

            for col in range(1, max_cols + 1):

                val = ws.cell(row=fila_destino, column=col).value

                if val is not None:

                    val_str = str(val).strip()

                    if (

                        val_str != "" 

                        and val_str != "0" 

                        and val_str != "0.0" 

                        and not val_str.startswith("=")

                    ):

                        vacia = False

                        break 

            if vacia:

                break 

            fila_destino += 1



        # ====================================================

        # CONSERVAR EL ESTILO DE LA FILA ANTERIOR
        # Copia fuente, relleno, bordes, alineación y formato numérico,
        # sin copiar valores ni fórmulas de la fila modelo.
        fila_modelo = fila_destino - 1
        if fila_modelo >= fila_encabezados + 1:
            for col in range(1, ws.max_column + 1):
                celda_modelo = ws.cell(row=fila_modelo, column=col)
                celda_nueva = ws.cell(row=fila_destino, column=col)
                if isinstance(celda_nueva, MergedCell):
                    continue
                if celda_modelo.has_style:
                    celda_nueva._style = copy(celda_modelo._style)
                if celda_modelo.number_format:
                    celda_nueva.number_format = celda_modelo.number_format

            dimension_modelo = ws.row_dimensions[fila_modelo]
            dimension_nueva = ws.row_dimensions[fila_destino]
            if dimension_modelo.height is not None:
                dimension_nueva.height = dimension_modelo.height

        # ====================================================

        # FUNCIONES DE CONVERSIÓN Y ESCRITURA SEGURA

        # ====================================================



        def intentar_convertir_a_numero(texto):

            """Intenta pasar un string a número flotante o entero"""

            if not texto:

                return texto, False

            texto_str = str(texto).strip()

            if texto_str.startswith("0") and not texto_str.startswith("0.") and len(texto_str) > 1:

                return texto_str, False

            texto_limpio = texto_str.replace(",", "")

            try:

                if "." in texto_limpio:

                    return float(texto_limpio), True

                else:

                    return int(texto_limpio), True

            except ValueError:

                return texto_str, False



        def colocar(encabezado, valor_colocar, ocurrencia=1):



            columna = obtener_columna(encabezado, ocurrencia)



            if columna:



                celda_target = ws.cell(row=fila_destino, column=columna)



                for merged_range in ws.merged_cells.ranges:

                    if celda_target.coordinate in merged_range:

                        celda_target = ws.cell(

                            row=merged_range.min_row,

                            column=merged_range.min_col

                        )

                        break



                valor_final, _es_numero = intentar_convertir_a_numero(valor_colocar)
                # Se modifica únicamente el valor; se conserva el estilo de plantilla.
                celda_target.value = valor_final



        # ====================================================

        # ASIGNACIÓN DE CAMPOS

        # ====================================================



        colocar("LICENCIA", licencia)



        colocar("EXPEDIENTE", expediente_valor, 1)



        colocar("PROPIETARIO", propietario, 1)



        colocar("DIRECCION DEL INMUEBLE", direccion_excel)



        colocar("DIRECCION PARA NOTIFICAR", direccion_excel)



        colocar("FINCA", finca)



        colocar("FOLIO", folio)



        colocar("LIBRO", libro)



        colOCAR = colocar("CODIGO MUNICIPAL", codigo_municipal)



        colocar("AREA DEL PREDIO", area_del_terreno)



        colocar("AREA A CONSTRUIR", area_a_construir)



        colocar("MURO", muro)



        colocar("PLANIFICADOR", planificador)

        # Primera columna COLEGIADO de la fila 2 de BASE DE DATOS.
        colocar("COLEGIADO", colegiado_planificado, 1)



        if tipo_planificador == "ING":



            colocar("ING", "X", 1)



            colocar("ARQ", "-", 1)



        elif tipo_planificador == "ARQ":



            colocar("ING", "-", 1)



            colocar("ARQ", "X", 1)



        colocar("EJECUTOR", ejecutor)

        # Copiar el mismo colegiado del planificador a la segunda columna
        # COLEGIADO, correspondiente al ejecutor.
        colocar("COLEGIADO", colegiado_planificado, 2)



        if tipo_ejecutor == "ING":



            colocar("ING", "X", 2)



            colocar("ARQ", "-", 2)



        elif tipo_ejecutor == "ARQ":



            colocar("ING", "-", 2)



            colocar("ARQ", "X", 2)



        colocar("GESTOR", gestor)



        colocar("DESTINO DE LA OBRA", destino_obra)



        colocar("CARACTERISTICAS 1", nombre_proyecto)



        colocar("CARACTERISTICAS 2", valor(datos, "observaciones"))



        colocar("UNIDADES FUNCIONALES", unidades_funcionales)



        colocar("NIVELES", niveles)



        colocar("COSTO DE OBRA", costo_obra)



        colocar("EXPEDIENTE", expediente_valor, 2)



        colocar("PROPIETARIO", propietario, 2)



        colocar("PAGO DE LICENCIA", valor(datos, "V_LICENCIA"))



        colocar("RECIBO", valor(datos, "RECIBO_LICENCIA_OBRA"), 1)



        colocar("PAGO DE ALINEACION", pago_alineacion)



        colocar("RECIBO", recibo_alineacion, 2)



        colocar("TIPO DE LICENCIA", valor(datos, "nombre_tramite"))



        colocar("NUMERO DE LICENCIA", licencia)



        colocar("NUMERO DE SUPERVISIONES", visitas_minimas)



        colocar("FECHA DE PAGO", fecha_pago)



        colocar("VENCE", se_vence)



        colocar("LICENCIA DE USO DE SUELO", numero_licencia_suelo)



        colocar("RESOLUCION DEL MARN", valor(datos, "comentarios"))



        if "RESIDENCIAL" in solicita_mayus:



            colocar("AREA RESIDENCIAL", area_a_construir)



            colocar("USO RESIDENCIAL", area_a_construir)



        colocar("CLASIFICACION DE USO", solicita)



        colocar("PAGO USO DE SUELO", pago_uso_suelo)



        colocar("RECIBO", recibo_uso_suelo, 3)



        colocar("FECHA DE VENCIMIENTO", se_vence)



        # ====================================================

        # GUARDAR CAMBIOS EN EL ARCHIVO

        # ====================================================



        try:
            wb.save(excel_path)
        finally:
            wb.close()
        # El historial solo se registra DESPUÉS de guardar el Excel con éxito.
        return {
            "expediente": expediente_valor,
            "propietario": propietario,
            "direccion": direccion_excel,
            "zona": zona,
            "tramite": valor(datos, "nombre_tramite"),
            "licencia": licencia,
            "licencia_suelo": numero_licencia_suelo,
            "planificador": planificador,
            "ejecutor": ejecutor,
            "colegiado": colegiado_planificado,
            "archivo_excel": os.path.abspath(excel_path),
            "hoja": ws.title,
            "fila_excel": fila_destino,
        }





# ============================================================

    def closeEvent(self, evento):
        self._cerrando = True
        self.temporizador_conexion.stop()
        return super().closeEvent(evento)


# EJECUTAR

# ============================================================




if __name__ == "__main__":
    app = QApplication(sys.argv)
    # Nombre estable para los datos locales incluso tras cambiar de versión.
    app.setOrganizationName("ControlDeObras")
    app.setApplicationName("DCO_Expedientes")

    # Mostrar la aplicación inmediatamente: el monitor NO bloquea el inicio.
    ventana = DCOApp()
    ventana.show()

    # Comprueba GitHub tras abrir la ventana y, después, cada 30 minutos.
    # Si está buscando o exportando, espera a terminar antes de avisar.
    QTimer.singleShot(1200, ventana.monitor_actualizaciones.iniciar)
    app.aboutToQuit.connect(ventana.monitor_actualizaciones.detener)

    sys.exit(app.exec())